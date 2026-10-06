"""
Importação de XLSX — levantamento de competências e adesão ao SIG.

Lê a planilha do diagnóstico (uma aba por secretaria: SEMUR, SEMFAZ, SEMUSB,
SEMA, SEMTRAN... com Departamento / Setor / Competência / respostas sobre o
SIG) e devolve os registros já limpos, resolvendo os problemas que esse tipo de
planilha sempre tem:

1. CÉLULAS MESCLADAS: "Departamento", "Setor", "Competência" e as respostas por
   setor (módulo, problema, observação...) ficam mescladas por várias linhas. O
   openpyxl devolve o valor só na primeira linha e None nas demais, então aqui o
   valor é propagado para baixo.
2. LOTAÇÃO SEM CABEÇALHO PRÓPRIO: o cabeçalho "setor" fica mesclado por cima de
   duas colunas (departamento + setor) — ou três, na SEMA, que tem uma coluna de
   coordenadoria antes. As colunas da lotação são descobertas por essa mescla.
3. LINHAS DE CONTINUAÇÃO: uma competência que usa várias camadas do SIG ocupa
   várias linhas; as de baixo trazem só a camada. Elas herdam a competência de
   cima e viram registros próprios (um por camada / forma de uso).
4. RESPONSÁVEL EMBUTIDO NO TEXTO: o nome e o cargo vêm dentro do próprio nome do
   setor — "Divisão de Controle e Informação - DICI (Gerente Damaris Dias)" ou
   "Divisão de Protocolo - Gerente Saulo". Aqui isso é separado em campos.

Além dos registros, exporta TODAS as abas da pasta de trabalho (inclusive as de
resumo, como "Estrutura interna") como tabelas, para a tela "Planilha" do app.

Uso:
    python scripts/importacao_de_xlsx.py planilha.xlsx
    python scripts/importacao_de_xlsx.py planilha.xlsx --aba "SEMUR"
    python scripts/importacao_de_xlsx.py planilha.xlsx --json saida.json
    python scripts/importacao_de_xlsx.py planilha.xlsx --ts src/data/registros.gerado.ts
    python scripts/importacao_de_xlsx.py planilha.xlsx --abas-ts src/data/planilha.gerado.ts
"""

from __future__ import annotations

import argparse
import datetime
import json
import re
import sys
import unicodedata
from dataclasses import dataclass, asdict, field
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    print("Falta a biblioteca openpyxl. Instale com:  pip install openpyxl")
    raise SystemExit(1)


# ---------------------------------------------------------------------------
# Modelo
# ---------------------------------------------------------------------------


@dataclass
class Registro:
    """Uma linha do levantamento: uma competência de um setor."""

    id: str

    # Secretaria = nome da aba da planilha (SEMUR, SEMFAZ, ...)
    secretaria: str = ""

    # Lotação
    departamento: str = ""
    departamentoSigla: str = ""
    departamentoResponsavel: str = ""
    departamentoCargo: str = ""

    setor: str = ""
    setorSigla: str = ""
    responsavel: str = ""
    cargo: str = ""

    # Competência avaliada
    competencia: str = ""

    # Respostas sobre o SIG
    sistemaAtende: bool | None = None
    comoAtende: str = ""
    estruturadoNoSistema: bool | None = None
    houveTreinamento: bool | None = None

    informacoesAcessadas: str = ""
    nivelAcessoCamada: str = ""
    acessoCamadaConcedido: bool | None = None
    iniciouAtividade: bool | None = None

    moduloSig: str = ""
    nivelAcessoModulo: str = ""
    acessoModuloConcedido: bool | None = None

    temProblema: bool | None = None
    temChamado: bool | None = None
    detalheProblema: str = ""
    observacao: str = ""

    # Preenchido depois, se houver planilha de contatos
    contatoInstitucional: dict[str, str] = field(default_factory=dict)
    contatoPessoal: dict[str, str] = field(default_factory=dict)


# ---------------------------------------------------------------------------
# Normalização
# ---------------------------------------------------------------------------

# Resultados de fórmula quebrada no Excel: não são dado, são lixo de cálculo.
VALORES_DE_ERRO = {"#DIV/0!", "#VALUE!", "#NAME?", "#REF!", "#N/A", "#NUM!", "#NULL!"}


def _chave(texto: Any) -> str:
    """Reduz um cabeçalho a letras e números, sem acento — para comparação."""
    if texto is None:
        return ""
    sem_acento = unicodedata.normalize("NFD", str(texto))
    sem_acento = "".join(c for c in sem_acento if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]", "", sem_acento.lower())


def _texto(valor: Any) -> str:
    """Converte a célula em texto limpo. '-', vazios e erros viram string vazia."""
    if valor is None:
        return ""
    texto = re.sub(r"\s+", " ", str(valor)).strip()
    if texto in VALORES_DE_ERRO:
        return ""
    return "" if texto in {"-", "--", "—", "N/A", "n/a"} else texto


def _booleano(valor: Any) -> bool | None:
    """Sim/Não → True/False. Qualquer outra coisa (inclusive '-') vira None."""
    chave = _chave(valor)
    if chave in {"sim", "s", "true", "1"}:
        return True
    if chave in {"nao", "n", "false", "0"}:
        return False
    return None


# ---------------------------------------------------------------------------
# Cabeçalhos aceitos
# ---------------------------------------------------------------------------

# A comparação é feita por "começa com" sobre a chave normalizada, porque os
# cabeçalhos da planilha são frases longas que podem variar na ponta.
COLUNAS: dict[str, list[str]] = {
    "departamento": ["departamento"],
    "setor": ["setor", "divisao", "unidade"],
    "competencia": ["competencia", "competenciaconformeregimento", "atribuicao"],
    "sistemaAtende": [
        "1osistemaatende",
        "1sistemaatende",
        # A planilha real escreve "sitema" (sem o s) — sem esta grafia a coluna
        # mais importante do levantamento ficaria de fora.
        "1ositemaatende",
        "1sitemaatende",
        "osistemaatendeaatividade",
        "sistemaatendeaatividade",
    ],
    "comoAtende": [
        "2osistemaatende",
        "2sistemaatende",
        "2ositemaatende",
        "osistemaatendedequalforma",
        "sistemaatendedequalforma",
    ],
    "estruturadoNoSistema": [
        "3jaestaestruturado",
        "3ojaestaestruturado",
        "jaestaestruturado",
        "estruturadonosistema",
    ],
    "houveTreinamento": ["houvetreinamento"],
    "informacoesAcessadas": ["informacoesaseremacessadas", "informacoesacessadas"],
    "nivelAcessoCamada": ["niveldeacessoacamada", "nivelacessocamada"],
    "acessoCamadaConcedido": ["niveldeacessoacamadafoiconcedido", "acessocamadaconcedido"],
    "iniciouAtividade": ["iniciouaatividade", "iniciouatividade"],
    "moduloSig": ["modulodosig", "modulosig", "modulo"],
    "nivelAcessoModulo": ["niveldeacessoaomodulo", "nivelacessomodulo"],
    "acessoModuloConcedido": ["niveldeacessoaomoduloconcedido", "acessomoduloconcedido"],
    "temProblema": ["temalgumproblema", "temproblema"],
    "temChamado": ["temchamadoaberto", "temchamado"],
    "detalheProblema": ["detalheoproblema", "detalheproblema"],
    "observacao": ["observacao", "observacoes"],
}

# Campos que devem ser lidos como Sim/Não.
CAMPOS_BOOLEANOS = {
    "sistemaAtende",
    "estruturadoNoSistema",
    "houveTreinamento",
    "acessoCamadaConcedido",
    "iniciouAtividade",
    "acessoModuloConcedido",
    "temProblema",
    "temChamado",
}

# Campos da competência em si: uma linha de continuação (só com outra camada)
# herda estes da linha de cima quando vierem vazios.
CAMPOS_DA_COMPETENCIA = {"sistemaAtende", "comoAtende"}


def _mapear_colunas(cabecalho: list[Any]) -> dict[str, int]:
    """Descobre em qual coluna está cada campo, comparando por prefixo."""
    chaves = [_chave(celula) for celula in cabecalho]
    mapa: dict[str, int] = {}

    for campo, aceitos in COLUNAS.items():
        for indice, chave in enumerate(chaves):
            if not chave or indice in mapa.values():
                continue
            if any(chave.startswith(aceito) or aceito.startswith(chave) for aceito in aceitos if chave):
                mapa[campo] = indice
                break

    # O cabeçalho "setor" mesclado por cima de várias colunas (já desmesclado,
    # aparece repetido): a última é o setor, a penúltima o departamento. Uma
    # terceira, à esquerda (SEMA), é a coordenadoria — fica só na aba completa.
    lotacao = [i for i, chave in enumerate(chaves) if chave == "setor"]
    if len(lotacao) >= 2:
        mapa["setor"] = lotacao[-1]
        mapa.setdefault("departamento", lotacao[-2])
        if mapa["departamento"] == mapa["setor"]:
            mapa["departamento"] = lotacao[-2]

    return mapa


# ---------------------------------------------------------------------------
# Extração de responsável ("... - DICI (Gerente Fulana de Tal)")
# ---------------------------------------------------------------------------

CARGOS_CONHECIDOS = [
    "Diretor",
    "Diretora",
    "Gerente",
    "Chefe",
    "Coordenador",
    "Coordenadora",
    "Secretário",
    "Secretária",
    "Responsável",
]

# Do mais longo para o mais curto: evita que "Diretor" vença "Diretora".
CARGOS_CONHECIDOS.sort(key=len, reverse=True)


def _separar_cargo(conteudo: str) -> tuple[str, str]:
    """'Gerente Fulana de Tal' → ('Fulana de Tal', 'Gerente'). Sem cargo → ('', '')."""
    # Compara a PRIMEIRA PALAVRA para nao confundir "Diretora" com "Diretor"
    # (o prefixo curto casaria e sobraria um "a" grudado no nome).
    primeira_palavra = conteudo.split(" ")[0] if conteudo else ""
    for possivel in CARGOS_CONHECIDOS:
        if _chave(primeira_palavra) == _chave(possivel):
            return _texto(conteudo[len(primeira_palavra) :]), possivel
    return "", ""


PALAVRAS_DE_UNIDADE = {"competencia", "departamento", "divisao", "setor", "secretaria", "gerencia"}


def _parece_nota(conteudo: str) -> bool:
    """Texto entre parênteses que fala da unidade (e não é o nome de alguém)."""
    return any(_chave(palavra) in PALAVRAS_DE_UNIDADE for palavra in conteudo.split())


def separar_unidade(texto_bruto: str) -> tuple[str, str, str, str]:
    """
    Quebra "Divisão de Controle e Informação - DICI (Gerente Damaris Dias)" em:
        (nome_da_unidade, sigla, nome_do_responsavel, cargo)

    Se algum pedaço não existir, volta string vazia — nunca levanta erro.
    """
    texto = _texto(texto_bruto)
    if not texto:
        return "", "", "", ""

    responsavel = ""
    cargo = ""

    # O responsável vem entre parênteses, geralmente precedido do cargo.
    parenteses = re.search(r"\(([^)]*)\)", texto)
    if parenteses:
        conteudo = _texto(parenteses.group(1))
        texto = _texto(texto[: parenteses.start()] + texto[parenteses.end() :])
        responsavel, cargo = _separar_cargo(conteudo)
        if not cargo and _parece_nota(conteudo):
            # "(Competência do Departamento)" é nota sobre a unidade, não pessoa.
            texto = _texto(texto_bruto)
        elif not cargo:
            responsavel = conteudo

    # A sigla vem depois de um hífen, em caixa alta. Às vezes, no lugar dela,
    # vem o responsável: "Divisão de Protocolo - Gerente Saulo".
    sigla = ""
    hifen = re.split(r"\s[-–]\s", texto)
    if len(hifen) > 1:
        candidata = _texto(hifen[-1])
        nome_apos_hifen, cargo_apos_hifen = _separar_cargo(candidata)
        if cargo_apos_hifen and not cargo:
            responsavel, cargo = nome_apos_hifen, cargo_apos_hifen
            texto = _texto(" - ".join(hifen[:-1]))
        elif candidata and candidata.upper() == candidata and len(candidata) <= 12:
            sigla = candidata
            texto = _texto(" - ".join(hifen[:-1]))

    return texto.rstrip(" :"), sigla, responsavel, cargo


# ---------------------------------------------------------------------------
# Leitura da planilha
# ---------------------------------------------------------------------------


def _desmesclar(planilha) -> None:
    """
    Copia o valor da célula mesclada para todas as células do intervalo.

    Sem isso, "Departamento" e "Setor" apareceriam só na primeira linha do
    bloco e todas as competências seguintes ficariam órfãs.
    """
    for intervalo in list(planilha.merged_cells.ranges):
        min_col, min_lin, max_col, max_lin = intervalo.bounds
        valor = planilha.cell(row=min_lin, column=min_col).value
        planilha.unmerge_cells(str(intervalo))

        for linha in range(min_lin, max_lin + 1):
            for coluna in range(min_col, max_col + 1):
                planilha.cell(row=linha, column=coluna).value = valor


def _encontrar_linha_cabecalho(linhas: list[tuple], limite: int = 10) -> int:
    """Acha a linha que contém os cabeçalhos (a que mais casa com COLUNAS)."""
    melhor_linha, melhor_placar = 0, 0

    for indice, linha in enumerate(linhas[:limite]):
        placar = len(_mapear_colunas(list(linha)))
        if placar > melhor_placar:
            melhor_linha, melhor_placar = indice, placar

    return melhor_linha


def _slug(texto: str) -> str:
    return _chave(texto) or "aba"


def _importar_aba(planilha) -> list[Registro]:
    """Lê uma aba de levantamento. Aba sem coluna de competência devolve []."""
    _desmesclar(planilha)

    linhas = list(planilha.iter_rows(values_only=True))
    if not linhas:
        return []

    indice_cabecalho = _encontrar_linha_cabecalho(linhas)
    mapa = _mapear_colunas(list(linhas[indice_cabecalho]))

    if "competencia" not in mapa:
        return []

    secretaria = _texto(planilha.title)
    prefixo = _slug(secretaria)

    # Colunas que são resposta (não lotação nem competência): usadas para saber
    # se uma linha sem competência é continuação da de cima ou só lixo.
    colunas_lotacao = {i for i, c in enumerate(linhas[indice_cabecalho]) if _chave(c) == "setor"}
    colunas_lotacao |= {mapa.get("departamento", -1), mapa.get("setor", -1)}
    colunas_resposta = [
        indice
        for campo, indice in mapa.items()
        if campo not in {"departamento", "setor", "competencia"} and indice not in colunas_lotacao
    ]

    registros: list[Registro] = []
    ultimo_departamento = ""
    ultimo_setor = ""
    anterior: Registro | None = None

    for numero, linha in enumerate(linhas[indice_cabecalho + 1 :], start=indice_cabecalho + 2):

        def celula(campo: str) -> Any:
            indice = mapa.get(campo)
            return linha[indice] if indice is not None and indice < len(linha) else None

        competencia = _texto(celula("competencia"))

        # Propaga a lotação para baixo (planilhas costumam repetir só na 1a linha).
        departamento_bruto = _texto(celula("departamento")) or ultimo_departamento
        # Departamento novo sem setor preenchido (ex.: "Departamento Tributário"
        # listado sem divisões): o setor de cima é de OUTRO departamento.
        if departamento_bruto != ultimo_departamento:
            ultimo_setor = ""
            anterior = None
        setor_bruto = _texto(celula("setor")) or ultimo_setor
        mudou_setor = setor_bruto != ultimo_setor
        ultimo_departamento = departamento_bruto
        ultimo_setor = setor_bruto

        tem_resposta = any(
            _texto(linha[i]) for i in colunas_resposta if i < len(linha)
        )

        # Linha de continuação: sem competência, mas com outra camada / forma de
        # uso, no mesmo setor da competência de cima.
        continuacao = not competencia and tem_resposta and anterior is not None and not mudou_setor
        if continuacao:
            competencia = anterior.competencia

        if not competencia:
            continue

        dep_nome, dep_sigla, dep_resp, dep_cargo = separar_unidade(departamento_bruto)
        set_nome, set_sigla, set_resp, set_cargo = separar_unidade(setor_bruto)

        registro = Registro(
            id=f"{prefixo}-{numero:04d}",
            secretaria=secretaria,
            departamento=dep_nome,
            departamentoSigla=dep_sigla,
            departamentoResponsavel=dep_resp,
            departamentoCargo=dep_cargo,
            setor=set_nome,
            setorSigla=set_sigla,
            responsavel=set_resp,
            cargo=set_cargo,
            competencia=competencia,
        )

        for campo in COLUNAS:
            if campo in {"departamento", "setor", "competencia"}:
                continue

            valor = celula(campo)
            if campo in CAMPOS_BOOLEANOS:
                setattr(registro, campo, _booleano(valor))
            else:
                setattr(registro, campo, _texto(valor))

            if continuacao and campo in CAMPOS_DA_COMPETENCIA and getattr(registro, campo) in (None, ""):
                setattr(registro, campo, getattr(anterior, campo))

        registros.append(registro)
        anterior = registro

    return registros


def importacao_de_xlsx(caminho: str | Path, aba: str | None = None) -> list[Registro]:
    """
    Importa a planilha de levantamento e devolve a lista de registros.

    :param caminho: caminho do arquivo .xlsx
    :param aba: nome da aba; se omitido, lê TODAS as abas que tiverem a coluna
                de competência (as de resumo são ignoradas aqui)
    """
    arquivo = Path(caminho)
    if not arquivo.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {arquivo}")

    livro = load_workbook(arquivo, data_only=True)
    abas = [livro[aba]] if aba else livro.worksheets

    registros: list[Registro] = []
    for planilha in abas:
        registros.extend(_importar_aba(planilha))

    if aba and not registros:
        raise ValueError(f"Não encontrei a coluna de competência na aba '{aba}'.")

    return registros


# ---------------------------------------------------------------------------
# Todas as abas, como tabelas (tela "Planilha" do app)
# ---------------------------------------------------------------------------

# Nomes das colunas de lotação quando o cabeçalho "setor" cobre várias colunas.
NIVEIS_LOTACAO = ["Setor", "Departamento", "Coordenadoria"]


def _formatar(valor: Any, porcentagem: bool) -> str:
    """Célula → texto para exibir. Números perdem as dízimas do Excel."""
    if isinstance(valor, bool):
        return "Sim" if valor else "Não"
    if isinstance(valor, (int, float)):
        if porcentagem:
            return f"{valor:.1f}".replace(".", ",") + "%"
        if float(valor).is_integer():
            return str(int(valor))
        return f"{valor:.2f}".rstrip("0").rstrip(".").replace(".", ",")
    if isinstance(valor, (datetime.datetime, datetime.date)):
        return valor.strftime("%d/%m/%Y")
    return _texto(valor)


def exportar_abas(caminho: str | Path) -> list[dict[str, Any]]:
    """
    Devolve TODAS as abas da pasta de trabalho como tabelas:
        {nome, colunas, colunasLotacao, linhas: [{linha, valores}]}

    Mescla vertical é repetida para baixo (cada linha fica legível sozinha);
    mescla horizontal fica só na primeira coluna (senão o departamento
    apareceria duplicado também na coluna do setor). Linha que só teria
    valores herdados de mescla é descartada — ela não traz nada de novo.
    """
    livro = load_workbook(Path(caminho), data_only=True)
    abas: list[dict[str, Any]] = []

    for planilha in livro.worksheets:
        max_lin, max_col = planilha.max_row, planilha.max_column
        celulas = {
            (lin, col): planilha.cell(row=lin, column=col).value
            for lin in range(1, max_lin + 1)
            for col in range(1, max_col + 1)
        }
        proprias = {pos for pos, valor in celulas.items() if _texto(valor)}

        # Cabeçalho = primeira linha com conteúdo.
        linha_cabecalho = min((lin for lin, _ in proprias), default=1)

        mesclas_cabecalho: dict[int, tuple[int, int]] = {}
        for intervalo in planilha.merged_cells.ranges:
            min_col, min_lin, max_col_m, max_lin_m = intervalo.bounds
            valor = celulas.get((min_lin, min_col))
            if min_lin <= linha_cabecalho <= max_lin_m:
                for col in range(min_col, max_col_m + 1):
                    mesclas_cabecalho[col] = (min_col, max_col_m)
            for lin in range(min_lin, max_lin_m + 1):
                if lin != min_lin:
                    celulas[(lin, min_col)] = valor
                for col in range(min_col + 1, max_col_m + 1):
                    celulas[(lin, col)] = None

        # Colunas usadas: as que têm qualquer valor (próprio) na aba.
        usadas = sorted({col for _, col in proprias})

        colunas: list[str] = []
        lotacao: list[int] = []
        for posicao, col in enumerate(usadas):
            inicio, fim = mesclas_cabecalho.get(col, (col, col))
            titulo = _texto(planilha.cell(row=linha_cabecalho, column=inicio).value)
            if _chave(titulo) == "setor":
                nivel = fim - col
                titulo = NIVEIS_LOTACAO[nivel] if nivel < len(NIVEIS_LOTACAO) else titulo
                lotacao.append(posicao)
            elif not titulo and posicao == 0:
                titulo = "Secretaria"
                lotacao.append(posicao)
            elif not titulo:
                titulo = f"Coluna {get_column_letter(col)}"
            elif inicio != col:
                titulo = f"{titulo} ({col - inicio + 1})"
            colunas.append(titulo)

        porcentagem = ["%" in c for c in colunas]

        linhas: list[dict[str, Any]] = []
        for lin in range(linha_cabecalho + 1, max_lin + 1):
            valores = [_formatar(celulas.get((lin, col)), porcentagem[i]) for i, col in enumerate(usadas)]
            tem_proprio = any((lin, col) in proprias and valores[i] for i, col in enumerate(usadas))
            if tem_proprio:
                linhas.append({"linha": lin, "valores": valores})

        abas.append(
            {
                "nome": _texto(planilha.title),
                "colunas": colunas if linhas else [],
                "colunasLotacao": lotacao if linhas else [],
                "linhas": linhas,
            }
        )

    return abas


# ---------------------------------------------------------------------------
# Saídas
# ---------------------------------------------------------------------------


def gerar_json(registros: list[Registro]) -> str:
    return json.dumps([asdict(r) for r in registros], ensure_ascii=False, indent=2)


def gerar_typescript(registros: list[Registro], origem: str) -> str:
    """Gera o arquivo que o app React Native consome."""
    corpo = json.dumps([asdict(r) for r in registros], ensure_ascii=False, indent=2)
    corpo = corpo.replace(": null", ": undefined")

    return (
        "/**\n"
        " * ARQUIVO GERADO AUTOMATICAMENTE — NAO EDITE A MAO.\n"
        " *\n"
        f" * Origem: {origem}\n"
        f" * Registros: {len(registros)}\n"
        f" * Regenerar: python scripts/importacao_de_xlsx.py dados/{origem} --ts src/data/registros.gerado.ts\n"
        " */\n\n"
        "import type { Registro } from './registros';\n\n"
        f"export const registrosImportados: Registro[] = {corpo};\n"
    )


def gerar_typescript_abas(abas: list[dict[str, Any]], origem: str) -> str:
    """Gera o arquivo com todas as abas da planilha (tela "Planilha")."""
    # Uma aba por linha, sem indentação: o arquivo fica bem menor no bundle.
    corpo = ",\n".join(json.dumps(aba, ensure_ascii=False) for aba in abas)

    return (
        "/**\n"
        " * ARQUIVO GERADO AUTOMATICAMENTE — NAO EDITE A MAO.\n"
        " *\n"
        f" * Origem: {origem}\n"
        f" * Abas: {len(abas)} ({', '.join(a['nome'] for a in abas)})\n"
        f" * Regenerar: python scripts/importacao_de_xlsx.py dados/{origem} --abas-ts src/data/planilha.gerado.ts\n"
        " */\n\n"
        "import type { AbaPlanilha } from './planilha';\n\n"
        f"export const origemPlanilha = {json.dumps(origem, ensure_ascii=False)};\n\n"
        f"export const abasPlanilha: AbaPlanilha[] = [\n{corpo}\n];\n"
    )


def _resumo(registros: list[Registro]) -> str:
    secretarias = sorted({r.secretaria for r in registros if r.secretaria})
    setores = sorted({(r.secretaria, r.setorSigla or r.setor) for r in registros if r.setorSigla or r.setor})
    atendidas = sum(1 for r in registros if r.sistemaAtende is True)
    nao_atendidas = sum(1 for r in registros if r.sistemaAtende is False)
    sem_resposta = sum(1 for r in registros if r.sistemaAtende is None)
    com_problema = sum(1 for r in registros if r.temProblema is True)
    sem_chamado = sum(1 for r in registros if r.temProblema is True and r.temChamado is False)

    linhas = [
        f"Registros importados....: {len(registros)}",
        f"Secretarias.............: {len(secretarias)} ({', '.join(secretarias) or '-'})",
        f"Setores.................: {len(setores)}",
        f"Sistema atende..........: {atendidas} sim / {nao_atendidas} não / {sem_resposta} sem resposta",
        f"Com problema no SIG.....: {com_problema}",
        f"Problema SEM chamado....: {sem_chamado}",
    ]
    for secretaria in secretarias:
        da_secretaria = [r for r in registros if r.secretaria == secretaria]
        linhas.append(f"  {secretaria:<10} {len(da_secretaria):>4} registros")
    return "\n".join(linhas)


def main() -> int:
    analisador = argparse.ArgumentParser(
        description="Importa a planilha XLSX do levantamento de competências / adesão ao SIG.",
    )
    analisador.add_argument("planilha", help="caminho do arquivo .xlsx")
    analisador.add_argument("--aba", help="nome da aba (padrão: todas as abas de levantamento)")
    analisador.add_argument("--json", help="grava a saída em JSON neste caminho")
    analisador.add_argument("--ts", help="grava o arquivo TypeScript do app neste caminho")
    analisador.add_argument(
        "--abas-ts", help="grava TODAS as abas (tabelas) no arquivo TypeScript da tela Planilha"
    )

    argumentos = analisador.parse_args()

    try:
        registros = importacao_de_xlsx(argumentos.planilha, argumentos.aba)
    except (FileNotFoundError, ValueError) as erro:
        print(f"Erro: {erro}")
        return 1

    if not registros:
        print("Nenhuma linha de competência encontrada na planilha.")
        return 1

    print(_resumo(registros))
    origem = Path(argumentos.planilha).name

    if argumentos.json:
        Path(argumentos.json).write_text(gerar_json(registros), encoding="utf-8")
        print(f"JSON gravado em: {argumentos.json}")

    if argumentos.ts:
        destino = Path(argumentos.ts)
        destino.parent.mkdir(parents=True, exist_ok=True)
        destino.write_text(gerar_typescript(registros, origem), encoding="utf-8")
        print(f"TypeScript gravado em: {destino}")

    if argumentos.abas_ts:
        abas = exportar_abas(argumentos.planilha)
        destino = Path(argumentos.abas_ts)
        destino.parent.mkdir(parents=True, exist_ok=True)
        destino.write_text(gerar_typescript_abas(abas, origem), encoding="utf-8")
        for aba in abas:
            print(f"  aba {aba['nome']!r}: {len(aba['linhas'])} linhas, {len(aba['colunas'])} colunas")
        print(f"Abas gravadas em: {destino}")

    if not argumentos.json and not argumentos.ts and not argumentos.abas_ts:
        print("\n(nenhuma saída pedida — use --json, --ts e/ou --abas-ts para gravar)")

    return 0


if __name__ == "__main__":
    sys.exit(main())
